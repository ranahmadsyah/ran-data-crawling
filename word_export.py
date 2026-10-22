"""Build an in-memory Word document from source API content."""

from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from pathlib import PurePosixPath
from urllib.parse import urljoin, urlsplit
import re

import requests
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps, UnidentifiedImageError

from integrity_api import CMS, SITES


MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def filename(page):
    title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", page.title).strip(" .")[:100]
    if not title:
        title = PurePosixPath(urlsplit(page.url).path).name or "integrity-content"
    return title + ".docx"


def image_bytes(url):
    """Images are read only from the same four sites or their shared CMS."""
    trusted = {urlsplit(CMS).hostname} | set(SITES) | {"www." + host for host in SITES}
    for _ in range(4):
        parsed = urlsplit(url)
        if (parsed.scheme != "https" or parsed.hostname not in trusted or parsed.username
                or parsed.password or parsed.port not in {None, 443}):
            raise ValueError("Sumber gambar tidak didukung")
        with requests.get(url, timeout=(5, 12), stream=True, allow_redirects=False) as response:
            if response.is_redirect:
                url = urljoin(url, response.headers["Location"])
                continue
            response.raise_for_status()
            chunks, size = [], 0
            for chunk in response.iter_content(65536):
                size += len(chunk)
                if size > 8 * 1024 * 1024:
                    raise ValueError("Gambar terlalu besar")
                chunks.append(chunk)
            with Image.open(BytesIO(b"".join(chunks))) as image:
                if image.width * image.height > 30_000_000:
                    raise ValueError("Resolusi gambar terlalu besar")
                image = ImageOps.exif_transpose(image)
                output = BytesIO()
                image.convert("RGBA" if "A" in image.getbands() else "RGB").save(output, format="PNG")
                return output.getvalue()
    raise ValueError("Terlalu banyak pengalihan gambar")


class Renderer:
    def __init__(self, document, page, image_loader):
        self.document, self.page, self.image_loader = document, page, image_loader
        self.warnings = []
        self.images = {}

    def inline(self, paragraph, node, parent=None, bold=False, italic=False):
        parent = parent if parent is not None else paragraph._p
        if isinstance(node, NavigableString):
            run = paragraph.add_run(re.sub(r"\s+", " ", str(node)))
            run.bold, run.italic = bold, italic
            if parent is not paragraph._p:
                parent.append(run._r)
            return
        if not isinstance(node, Tag) or node.name in {"script", "style", "noscript"}:
            return
        if node.name == "br":
            run = OxmlElement("w:r")
            run.append(OxmlElement("w:br"))
            parent.append(run)
            return
        if node.name == "img":
            src = urljoin(self.page.url, node.get("src") or node.get("data-src") or "")
            try:
                if src not in self.images:
                    self.images[src] = self.image_loader(src)
                data = self.images[src]
                if isinstance(data, Exception):
                    raise ValueError("Gambar tidak tersedia") from data
                with Image.open(BytesIO(data)) as image:
                    natural_width = image.width * 2.54 / 96
                    requested = node.get("width", "")
                    if str(requested).isdigit():
                        natural_width = min(natural_width, int(requested) * 2.54 / 96)
                    width = min(Cm(natural_width), Cm(15.5), Cm(16) * image.width / image.height)
                picture = paragraph.add_run().add_picture(BytesIO(data), width=int(width))
                picture._inline.docPr.set("descr", node.get("alt", ""))
            except (requests.RequestException, ValueError, OSError, UnidentifiedImageError,
                    Image.DecompressionBombError):
                self.warnings.append("Gambar tidak dapat disertakan: " + src)
            return
        if node.name == "a":
            href = urljoin(self.page.url, node.get("href", ""))
            if urlsplit(href).scheme in {"http", "https", "mailto", "tel"}:
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(
                    href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
                parent.append(link)
                for child in node.children:
                    self.inline(paragraph, child, link, bold, italic)
                return
        for child in node.children:
            self.inline(paragraph, child, parent, bold or node.name in {"b", "strong"},
                        italic or node.name in {"i", "em"})

    def list(self, node, target, depth=0):
        # A new numbering instance for every list preserves restarts and ol[start].
        numbering = self.document.part.numbering_part.element
        abstract_id = max([int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))] + [-1]) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, val in [("start", node.get("start", "1")),
                         ("numFmt", "decimal" if node.name == "ol" else "bullet"),
                         ("lvlText", "%1." if node.name == "ol" else "•")]:
            item = OxmlElement("w:" + tag)
            item.set(qn("w:val"), str(val))
            level.append(item)
        abstract.append(level)
        numbering.append(abstract)
        number = numbering.add_num(abstract_id)
        for item in node.find_all("li", recursive=False):
            p = target.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 + 0.5 * depth)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = number.numId
            for child in item.children:
                if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                    self.list(child, target, depth + 1)
                else:
                    self.inline(p, child)

    def block(self, node, target):
        if isinstance(node, NavigableString):
            if node.strip():
                target.add_paragraph(str(node).strip())
            return
        if not isinstance(node, Tag) or node.name in {"script", "style", "noscript"}:
            return
        if node.name in {"ul", "ol"}:
            self.list(node, target)
        elif node.name == "table":
            rows = node.find_all("tr")
            if not rows:
                return
            table = target.add_table(rows=0, cols=max(len(r.find_all(["td", "th"], recursive=False)) for r in rows))
            table.style = "Table Grid"
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row.find_all(["td", "th"], recursive=False)):
                    p = cell.paragraphs[0]
                    for child in value.children:
                        self.inline(p, child, bold=value.name == "th")
        elif node.name in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "img"}:
            if not node.get_text(strip=True) and node.name != "img" and not node.find("img"):
                return
            style = "Heading " + str(min(max(int(node.name[1]) - 1, 1), 3)) if re.fullmatch(r"h[1-6]", node.name) else None
            p = target.add_paragraph(style=style)
            self.inline(p, node)
        elif not node.find(["p", "div", "section", "ul", "ol", "table", "h1", "h2", "h3"]):
            p = target.add_paragraph()
            self.inline(p, node)
        else:
            for child in node.children:
                self.block(child, target)


def build_docx(page, image_loader=image_bytes):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "Tahoma")
    fonts.set(qn("w:cs"), "Tahoma")
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        style = document.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("14675B")
        style.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=2)
    table.style, table.autofit = "Table Grid", False
    table.columns[0].width, table.columns[1].width = Cm(4.5), Cm(12.5)
    for cell, text in zip(table.rows[0].cells, ["Elemen", "Isi"]):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for label, value in page.metadata:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_paragraph()
    document.add_heading(page.title or "Judul: Kosong", 0)
    renderer = Renderer(document, page, image_loader)
    soup = BeautifulSoup(page.content, "html.parser")
    # Several service pages contain many client logos. Download unique images concurrently.
    urls = list(dict.fromkeys(urljoin(page.url, img.get("src") or img.get("data-src") or "")
                             for img in soup.find_all("img")))
    def load(url):
        try:
            return url, image_loader(url)
        except Exception as exc:
            return url, exc
    if urls:
        with ThreadPoolExecutor(max_workers=6) as pool:
            renderer.images.update(pool.map(load, urls))
    if not soup.get_text(strip=True) and not soup.find("img"):
        document.add_paragraph("Konten: Kosong")
    for node in soup.children:
        renderer.block(node, document)
    output = BytesIO()
    document.save(output)
    return output.getvalue(), list(dict.fromkeys(renderer.warnings))
