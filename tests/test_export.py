from io import BytesIO
from pathlib import Path
from unittest import TestCase, mock
from zipfile import ZipFile

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from streamlit.testing.v1 import AppTest

from integrity_api import Client, ExportError, Page, fetch_page, parse_url
from word_export import build_docx

APP = str(Path(__file__).resolve().parents[1] / "streamlit_app.py")


def sample_page(**changes):
    values = dict(url="https://www.integrity-indonesia.com/id/blog/contoh/",
                  site="Integrity Indonesia", language="id", title="Judul asli",
                  content="<p>Isi artikel asli.</p>", seo={"title": "Title asli", "focuskw": ""})
    return Page(**(values | changes))


class ApiTests(TestCase):
    def test_four_domains_use_correct_api_and_language(self):
        for host, path, code, language in [
            ("integrity-indonesia.com", "/id/blog/contoh/", "id", "id"),
            ("integrity-malaysia.com", "/services/example/", "ms", "en"),
            ("integrity-thailand.com", "/th/services/example/", "th", "th"),
            ("integrity-asia.com", "/services/example/", "asia", "en"),
        ]:
            with self.subTest(host=host):
                url, _, _, actual_code, actual_language = parse_url(host + path)
                self.assertEqual((actual_code, actual_language), (code, language))
                self.assertTrue(url.startswith("https://www."))
                self.assertTrue(Client(code).endpoint.endswith("/" + code + "/graphql"))

    def test_invalid_or_foreign_urls_never_make_requests(self):
        with mock.patch("integrity_api.requests.post") as request:
            for url in ["", "https://example.com/", "https://integrity-asia.com.evil.com/",
                        "https://user:pass@www.integrity-asia.com/", "file:///etc/passwd",
                        "https://www.integrity-asia.com:1234/", "https://www.integrity-asia.com/../secret"]:
                with self.subTest(url=url), self.assertRaises(ExportError):
                    fetch_page(url)
            request.assert_not_called()

    def test_thai_url_uses_translated_node_not_english_content(self):
        english = {"__typename": "Page", "id": "en", "languageCode": "en",
                   "translations": [{"id": "thai", "uri": "/th/services/example/", "languageCode": "th"}]}
        thai = {"title": "ชื่อเรื่อง", "content": "<p>เนื้อหาไทย</p>", "seo": {"focuskw": ""}}
        with mock.patch.object(Client, "query", side_effect=[{"nodeByUri": english}, {"page": thai}]) as query:
            result = fetch_page("https://www.integrity-thailand.com/th/services/example/")
        self.assertEqual(query.call_args_list[1].args[1], {"id": "thai"})
        self.assertEqual(result.title, "ชื่อเรื่อง")
        self.assertIn("เนื้อหาไทย", result.content)

    def test_missing_translation_does_not_silently_use_english(self):
        english = {"__typename": "Page", "id": "en", "languageCode": "en", "translations": []}
        with mock.patch.object(Client, "query", side_effect=[{"nodeByUri": english},
                    {"posts": {"nodes": []}, "pages": {"nodes": [english]}}]):
            with self.assertRaises(ExportError):
                Client("th").locate("/th/services/example/", "th")

    def test_blog_slug_lookup_handles_different_permalink(self):
        post = {"__typename": "Post", "id": "id1", "slug": "contoh", "uri": "/contoh/", "languageCode": "id"}
        with mock.patch.object(Client, "query", side_effect=[{"nodeByUri": None},
                                {"posts": {"nodes": [post]}, "pages": {"nodes": []}}]):
            self.assertEqual(Client("id").locate("/id/blog/contoh/", "id")["id"], "id1")

    def test_page_with_same_slug_but_wrong_parent_is_rejected(self):
        page = {"__typename": "Page", "id": "wrong", "slug": "example", "uri": "/ads/example/"}
        with mock.patch.object(Client, "query", side_effect=[{"nodeByUri": None},
                                {"posts": {"nodes": []}, "pages": {"nodes": [page]}}]):
            with self.assertRaises(ExportError):
                Client("asia").locate("/services/example/", "en")

    def test_service_url_cannot_resolve_an_unrelated_blog_post(self):
        post = {"__typename": "Post", "id": "wrong", "slug": "example", "uri": "/blog/example/"}
        with mock.patch.object(Client, "query", side_effect=[{"nodeByUri": None},
                                {"posts": {"nodes": [post]}, "pages": {"nodes": []}}]):
            with self.assertRaises(ExportError):
                Client("asia").locate("/services/example/", "en")

    def test_blank_api_content_is_reported_as_empty(self):
        node = {"__typename": "Page", "id": "empty", "languageCode": "en"}
        with mock.patch.object(Client, "query", side_effect=[{"nodeByUri": node}, {"page": {"title": "Home", "content": "", "seo": {}}}]):
            page = fetch_page("https://www.integrity-malaysia.com/")
        self.assertTrue(any("Konten: Kosong" in warning for warning in page.warnings))
        self.assertEqual([v for _, v in page.metadata][0], "Kosong")

    def test_service_uses_custom_fields_instead_of_reporting_empty_body(self):
        node = {"title": "CMS title", "content": "", "seo": {"title": "SEO asli"},
                "integrityPageTitle": {"pageTitleCustomTitle": "Judul layanan", "pageTitleDescription": "Pengantar asli"},
                "customFields": [{"__typename": "ParagraphSection", "heading": "Tentang layanan", "content": "<p>Penjelasan asli.</p>"},
                                 {"__typename": "FaqSection", "sectionTitle": "FAQ", "faqs": [{"question": "Pertanyaan asli?", "answer": "Jawaban asli."}]}]}
        with mock.patch.object(Client, "locate", return_value={"id": "service"}), mock.patch.object(Client, "read", return_value=node):
            page = fetch_page("https://www.integrity-asia.com/services/example/")
        text = BeautifulSoup(page.content, "html.parser").get_text(" ")
        self.assertEqual(page.title, "Judul layanan")
        for expected in ["Pengantar asli", "Tentang layanan", "Penjelasan asli.", "Pertanyaan asli?", "Jawaban asli."]:
            self.assertIn(expected, text)
        self.assertEqual(page.warnings, [])


class WordTests(TestCase):
    def test_metadata_is_separate_and_no_keyword_is_invented(self):
        data, warnings = build_docx(sample_page())
        doc = Document(BytesIO(data))
        rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
        self.assertEqual(rows, [["Elemen", "Isi"], ["Title tag", "Title asli"],
                               ["Meta description", "Kosong"], ["URL", sample_page().url],
                               ["Keyword utama (focus keyword)", "Kosong"]])
        self.assertIn("Isi artikel asli.", [p.text for p in doc.paragraphs])
        self.assertEqual(warnings, [])

    def test_unicode_lists_tables_and_links_survive_export(self):
        content = ('<h2>หัวข้อ</h2><p>Teks <strong>tebal</strong> dan <em>miring</em> � '
                   '<a href="https://www.integrity-asia.com/">tautan</a>.</p>'
                   '<ol start="3"><li>Pertama<ul><li>Anak</li></ul></li><li>Kedua</li></ol>'
                   '<table><tr><th>Label</th><th>Nilai</th></tr><tr><td>Asli</td><td>100%</td></tr></table>')
        data, _ = build_docx(sample_page(content=content))
        doc = Document(BytesIO(data))
        text = " ".join(p.text for p in doc.paragraphs)
        for expected in ["หัวข้อ", "tebal", "miring", "�", "tautan", "Pertama", "Anak", "Kedua"]:
            self.assertIn(expected, text)
        self.assertEqual(doc.tables[1].cell(1, 1).text, "100%")
        self.assertTrue(any(run.bold for p in doc.paragraphs for run in p.runs))
        with ZipFile(BytesIO(data)) as archive:
            self.assertIn(b'https://www.integrity-asia.com/', archive.read('word/_rels/document.xml.rels'))
            self.assertIn(b'w:val="3"', archive.read('word/numbering.xml'))

    def test_image_is_embedded_and_failures_are_reported(self):
        stream = BytesIO()
        Image.new("RGB", (20, 10), "white").save(stream, "PNG")
        page = sample_page(content='<p>Konten.</p><p><img src="https://cms-corporate.integrity-asia.com/image.png"></p>')
        data, warnings = build_docx(page, image_loader=lambda url: stream.getvalue())
        self.assertEqual(len(Document(BytesIO(data)).inline_shapes), 1)
        self.assertEqual(warnings, [])
        def unavailable(url):
            raise ValueError("Gagal")
        data, warnings = build_docx(page, image_loader=unavailable)
        self.assertEqual(len(warnings), 1)
        self.assertIn("Konten.", [p.text for p in Document(BytesIO(data)).paragraphs])


class AppTests(TestCase):
    def test_generate_download_delete_confirmation_and_cancel(self):
        app = AppTest.from_file(APP).run()
        self.assertEqual(len(app.exception), 0)
        with mock.patch("integrity_api.fetch_page", return_value=sample_page()) as fetch:
            app.text_input(key="page_url").input(sample_page().url)
            app.button[0].click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertIn("result", app.session_state)
            self.assertEqual(len(app.get("download_button")), 1)
            app.button(key="clear_result").click().run()
            self.assertIn("result", app.session_state)
            app.button(key="cancel_delete").click().run()
            self.assertIn("result", app.session_state)
            app.button(key="clear_result").click().run()
            app.button(key="confirm_delete").click().run()
            self.assertNotIn("result", app.session_state)
            self.assertEqual(fetch.call_count, 1)

    def test_failed_export_clears_stale_download_and_sessions_are_isolated(self):
        app = AppTest.from_file(APP).run()
        with mock.patch("integrity_api.fetch_page", return_value=sample_page()):
            app.text_input[0].input(sample_page().url)
            app.button[0].click().run()
        another = AppTest.from_file(APP).run()
        self.assertNotIn("result", another.session_state)
        with mock.patch("integrity_api.fetch_page", side_effect=ExportError("URL tidak ditemukan")):
            app.text_input[0].input("https://example.com/")
            app.button[0].click().run()
        self.assertNotIn("result", app.session_state)
        self.assertEqual(len(app.get("download_button")), 0)
        self.assertEqual(app.error[0].value, "URL tidak ditemukan")
